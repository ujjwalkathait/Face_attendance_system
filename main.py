import os.path
import datetime
import pickle
import tkinter as tk
import pandas as pd
import cv2
from openpyxl import Workbook
from PIL import Image, ImageTk
from openpyxl import load_workbook
import face_recognition
import util
from test import test


class App:
    def __init__(self):
        self.main_window = tk.Tk()
        self.sections = ['Section A', 'Section B']
        self.main_window.title('Attendance System')
        self.main_window.geometry("1200x520+350+100")

        self.choose_section = util.get_text_label(self.main_window, 'Choose section :- ')
        self.choose_section.place(x = 100, y=100)

        self.section_A_main_window = util.get_button(self.main_window, 'Section A', 'gray',lambda: self.main_page(self.sections[0]))
        self.section_A_main_window.place(x=400, y=200)

        self.section_B_main_window = util.get_button(self.main_window, 'Section B', 'gray',lambda: self.main_page(self.sections[1]))
        self.section_B_main_window.place(x=400, y=300)

    def main_page(self, section):
        self.main_page_window = tk.Toplevel(self.main_window)
        self.main_page_window.geometry("1200x520+370+120")

        self.section_name = section

        self.login_button_main_window = util.get_button(self.main_page_window, 'login', 'green', self.login)
        self.login_button_main_window.place(x=750, y=300)

        self.register_new_user_button_main_window = util.get_button(self.main_page_window, 'register new user', 'gray', self.register_new_user, fg='black')
        self.register_new_user_button_main_window.place(x=750, y=400)

        self.webcam_label = util.get_img_label(self.main_page_window)
        self.webcam_label.place(x=10, y=0, width=700, height=500)

        self.add_webcam(self.webcam_label)

        self.db_dir = './db'
        if not os.path.exists(self.db_dir):
            os.mkdir(self.db_dir)

        self.log_path = './Btech/{}'.format(section)

    def add_webcam(self, label):
        if 'cap' not in self.__dict__:
            self.cap = cv2.VideoCapture(1)

        self._label = label
        self.process_webcam()
 
    def process_webcam(self):
        ret, frame = self.cap.read()

        self.most_recent_capture_arr = frame
        img_ = cv2.cvtColor(self.most_recent_capture_arr, cv2.COLOR_BGR2RGB)
        self.most_recent_capture_pil = Image.fromarray(img_)
        imgtk = ImageTk.PhotoImage(image=self.most_recent_capture_pil)
        self._label.imgtk = imgtk
        self._label.configure(image=imgtk)

        self._label.after(20, self.process_webcam)

    def login(self):

        label = test(
                image=self.most_recent_capture_arr,
                model_dir='E:/python/Face_Recognition2/models',
                device_id=0
                )

        if label == 1:

            text = util.recognize(self.most_recent_capture_arr, self.db_dir)     

            name = text.split('-')[0]
            student_id = text.split('-')[1]
            current = datetime.datetime.now()

            if name in ['unknown_person', 'no_persons_found']:
                util.msg_box('Oops...', 'Unknown user. Please register new user or try again.')
            else:
                current_date = current.strftime('%Y-%m-%d')
                excel_file_path = os.path.join(self.log_path, f"{current_date}.xlsx")
                if os.path.exists(excel_file_path):
                    df = pd.read_excel(excel_file_path)
                    print(df)
                    if int(student_id) in df["Student Id"].astype(int).values:
                        util.msg_box('Oops...', 'User already logged in.')
                    else:
                        workbook = load_workbook(excel_file_path)
                        sheet = workbook.active
                        sheet.append([name, student_id, current.time(), 'P'])
                        workbook.save(excel_file_path)
                        util.msg_box("Hello !","You have been marked present {} with id - {} !".format(name, student_id))
                else:
                    workbook = Workbook()
                    sheet = workbook.active
                    sheet.title = "{}".format(self.section_name)
                    sheet.append(["Name", "Student Id", "Time", "Status"])
                    sheet.append([name, student_id , current.time(), 'P'])
                    workbook.save(excel_file_path)
                    util.msg_box("Hello !","You have been marked present {} with id - {} !".format(name, student_id))

        else:
            util.msg_box('Hey, you are a spoofer!', 'You are fake !')


    def register_new_user(self):
        self.register_new_user_window = tk.Toplevel(self.main_window)
        self.register_new_user_window.geometry("1200x520+370+120")

        self.accept_button_register_new_user_window = util.get_button(self.register_new_user_window, 'Accept', 'green', self.accept_register_new_user)
        self.accept_button_register_new_user_window.place(x=750, y=300)

        self.try_again_button_register_new_user_window = util.get_button(self.register_new_user_window, 'Try again', 'red', self.try_again_register_new_user)
        self.try_again_button_register_new_user_window.place(x=750, y=400)

        self.capture_label = util.get_img_label(self.register_new_user_window)
        self.capture_label.place(x=10, y=0, width=700, height=500)

        self.add_img_to_label(self.capture_label)

        self.entry_text_register_new_user = util.get_entry_text(self.register_new_user_window)
        self.entry_text_register_new_user.place(x=750, y=150)

        self.text_label_register_new_user = util.get_text_label(self.register_new_user_window, 'Please, \ninput username:')
        self.text_label_register_new_user.place(x=750, y=70)

    def try_again_register_new_user(self):
        self.register_new_user_window.destroy()

    def add_img_to_label(self, label):
        imgtk = ImageTk.PhotoImage(image=self.most_recent_capture_pil)
        label.imgtk = imgtk
        label.configure(image=imgtk)

        self.register_new_user_capture = self.most_recent_capture_arr.copy()

    def start(self):
        self.main_window.mainloop()

    def accept_register_new_user(self):
        text = self.entry_text_register_new_user.get(1.0, "end-1c")

        name = text.split('-')[0]
        student_id = text.split('-')[1]
        embeddings = face_recognition.face_encodings(self.register_new_user_capture)[0]

        file_path = os.path.join(self.db_dir, '{}-{}.pickle'.format(name,student_id))

        if os.path.exists(file_path):
            util.msg_box("Already Registered !", "You have already registered...")
        else:
            file = open(file_path, 'wb')
            pickle.dump(embeddings, file)
            util.msg_box('Success!', 'User was registered successfully !')

        self.register_new_user_window.destroy()


if __name__ == "__main__":
    app = App()
    app.start()







# import os
# import datetime
# import pickle

# import streamlit as st
# import pandas as pd
# import cv2
# import face_recognition
# from openpyxl import Workbook, load_workbook
# from PIL import Image

# # Utility functions
# import util
# from test import test

# def main():
#     st.title("Attendance System")
#     sections = ["Section A", "Section B"]
    
#     # Choose section
#     section = st.selectbox("Choose section", sections)
    
#     # Webcam and user options
#     frame_placeholder = st.empty()
#     options = st.radio("Choose an option", ["Login", "Register New User"])

#     # Webcam setup
#     cap = cv2.VideoCapture(1)

#     def process_frame():
#         ret, frame = cap.read()
#         frame_rgb = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
#         img_pil = Image.fromarray(frame_rgb)
#         return frame, img_pil

#     def display_frame():
#         _, img_pil = process_frame()
#         frame_placeholder.image(img_pil, caption="Webcam", use_column_width=True)

#     if options == "Login":
#         st.header("Login")
#         login_button = st.button("Login")

#         if login_button:
#             frame, _ = process_frame()
#             label = test(
#                 image=frame,
#                 model_dir='E:/python/Face_Recognition2/models',
#                 device_id=0
#             )

#             if label == 1:
#                 text = util.recognize(frame, './db')
#                 name = text.split('-')[0]
#                 student_id = text.split('-')[1]
#                 current = datetime.datetime.now()

#                 if name in ['unknown_person', 'no_persons_found']:
#                     st.warning("Unknown user. Please register or try again.")
#                 else:
#                     current_date = current.strftime('%Y-%m-%d')
#                     excel_file_path = f'./Btech/{section}/{current_date}.xlsx'

#                     if os.path.exists(excel_file_path):
#                         df = pd.read_excel(excel_file_path)
#                         if int(student_id) in df["Student Id"].astype(int).values:
#                             st.warning("User already logged in.")
#                         else:
#                             workbook = load_workbook(excel_file_path)
#                             sheet = workbook.active
#                             sheet.append([name, student_id, current.time(), 'P'])
#                             workbook.save(excel_file_path)
#                             st.success(f"You have been marked present: {name} with ID - {student_id}")
#                     else:
#                         workbook = Workbook()
#                         sheet = workbook.active
#                         sheet.title = section
#                         sheet.append(["Name", "Student Id", "Time", "Status"])
#                         sheet.append([name, student_id, current.time(), 'P'])
#                         workbook.save(excel_file_path)
#                         st.success(f"You have been marked present: {name} with ID - {student_id}")
#             else:
#                 st.error("You are a spoofer!")

#     elif options == "Register New User":
#         st.header("Register New User")
#         username = st.text_input("Enter username (format: Name-ID)")
#         register_button = st.button("Register")

#         if register_button and username:
#             frame, _ = process_frame()
#             try:
#                 name, student_id = username.split('-')
#                 embeddings = face_recognition.face_encodings(frame)[0]
#                 file_path = os.path.join('./db', f'{name}-{student_id}.pickle')

#                 if os.path.exists(file_path):
#                     st.warning("User already registered.")
#                 else:
#                     with open(file_path, 'wb') as f:
#                         pickle.dump(embeddings, f)
#                     st.success("User registered successfully!")
#             except Exception as e:
#                 st.error(f"Error: {e}")

#     display_frame()

# if __name__ == "__main__":
#     main()
